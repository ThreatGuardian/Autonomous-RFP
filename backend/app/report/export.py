"""Export the editable report to PDF (bid-report design) and Word (.docx)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    BaseDocTemplate, CondPageBreak, Frame, KeepTogether, PageBreak, PageTemplate, Paragraph, Spacer, Table, TableStyle,
)

from app.services.pdf_renderer import NumberedCanvas, esc
from app.services.report_renderer import (
    BAD, CARD, CARD_LINE, GOLD, GOOD, MUTED, WARN, Dots, ScoreBar, Tiles, _card, _styles, spaced,
)

TONES = {"good": GOOD, "warn": WARN, "bad": BAD, "neutral": MUTED}


def _visible(doc: dict[str, Any]) -> list[dict[str, Any]]:
    return [s for s in doc["sections"] if not s.get("hidden")]


def render_pdf(doc: dict[str, Any], path: Path, company: dict, approved: bool = False, *, label: str = "BID ANALYSIS",
               footer: str = "Bid analysis report", draft_note: str | None = "Draft for internal review") -> Path:
    st = _styles()
    page_w, page_h = A4
    margin = 22 * mm
    width = page_w - 2 * margin

    def decorate(c, _doc):
        c.saveState()
        x, y = margin, page_h - 15 * mm
        c.setFillColor(GOLD)
        p = c.beginPath()
        p.moveTo(x + 4, y + 8)
        p.lineTo(x + 8, y + 4)
        p.lineTo(x + 4, y)
        p.lineTo(x, y + 4)
        p.close()
        c.drawPath(p, stroke=0, fill=1)
        c.setFont("Inter-SemiBold", 7.4)
        c.drawString(x + 14, y + 1.5, " ".join(company["short_name"].upper()))
        c.setFont("Inter-Regular", 6.8)
        c.setFillColor(MUTED)
        c.drawRightString(page_w - margin, y + 7, label)
        c.setFont("Inter-SemiBold", 9.5)
        c.setFillColor(GOLD)
        c.drawRightString(page_w - margin, y - 4, " ".join(doc.get("quote_number", "")))
        c.setFont("Inter-Regular", 7)
        c.setFillColor(MUTED)
        c.drawString(margin, 10 * mm, f"{company['short_name']} · {footer}")
        if not approved and draft_note:
            c.drawCentredString(page_w / 2, 10 * mm, draft_note)
        c.restoreState()

    pdf = BaseDocTemplate(str(path), pagesize=A4, leftMargin=margin, rightMargin=margin, topMargin=28 * mm,
                          bottomMargin=22 * mm, title=doc["title"], author=company["name"])
    pdf.addPageTemplates([PageTemplate("r", [Frame(margin, 20 * mm, width, page_h - 48 * mm, id="f")], onPage=decorate)])
    s: list = [Paragraph(esc(doc["title"]), st["title"]), Spacer(1, 4), Paragraph(esc(doc.get("subtitle", "")), st["meta"]),
               Spacer(1, 10), Table([[""]], colWidths=[width], rowHeights=[1], style=[("LINEABOVE", (0, 0), (-1, -1), 0.6, CARD_LINE)])]
    cell = ParagraphStyle("tc", parent=st["body"], fontSize=8.2, leading=11.4)
    head = ParagraphStyle("th", parent=cell, fontName="Inter-SemiBold", textColor=MUTED)
    for sec in _visible(doc):
        flow: list = [Paragraph(spaced(sec["title"]), st["eyebrow"])]
        for b in sec["blocks"]:
            t = b["type"]
            if t == "lead":
                flow.append(Paragraph(esc(b["text"]), st["lead"]))
            elif t == "paragraph":
                flow += [Paragraph(esc(b["text"]), st["body"]), Spacer(1, 6)]
            elif t == "note":
                flow += [Paragraph(esc(b["text"]), st["small"]), Spacer(1, 4)]
            elif t == "bullets" and b.get("items"):
                flow += [Dots([(esc(i), GOLD) for i in b["items"]], width, st), Spacer(1, 4)]
            elif t == "kpis" and b.get("items"):
                flow += [Tiles([(i["value"], i["label"]) for i in b["items"][:4]], width, 48), Spacer(1, 8)]
            elif t == "bars":
                for i in b.get("items", []):
                    bar_label = i["label"] if len(i["label"]) < 46 else i["label"][:44] + "…"
                    flow.append(ScoreBar(bar_label, i.get("note", ""), float(i.get("value", 0)), width, str(i.get("display", ""))))
                flow.append(Spacer(1, 4))
            elif t == "table" and b.get("rows"):
                n = len(b["header"])
                data = [[Paragraph(esc(h), head) for h in b["header"]]] + [[Paragraph(esc(c), cell) for c in r] for r in b["rows"]]
                if b.get("widths"):
                    total = sum(b["widths"])
                    widths = [width * w / total for w in b["widths"]]
                else:
                    widths = [width / n] * n if n != 3 else [width * 0.22, width * 0.2, width * 0.58]
                tbl = Table(data, colWidths=widths, repeatRows=1)
                tbl.setStyle(TableStyle([
                    ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, -1), 0.5, CARD_LINE),
                    ("BACKGROUND", (0, 0), (-1, 0), CARD), ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4), ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ]))
                flow += [tbl, Spacer(1, 8)]
            elif t == "pagebreak":
                flow.append(PageBreak())
            elif t == "callout":
                tone = TONES.get(b.get("tone", "neutral"), MUTED)
                para = Paragraph(f"<font name='Inter-SemiBold' color='{tone.hexval()}'>{esc(b['text'])}</font>", st["body"])
                flow += [_card([para], width), Spacer(1, 8)]
        s.append(CondPageBreak(45 * mm))
        s += [KeepTogether(flow[:3]), *flow[3:]] if len(flow) > 3 else [KeepTogether(flow)]
    if approved:
        s += [Spacer(1, 10), Paragraph("Approved for submission.", st["small"])]
    pdf.build(s, canvasmaker=NumberedCanvas)
    return path


def render_docx(doc: dict[str, Any], path: Path, company: dict, *, label: str = "Bid analysis") -> Path:
    import docx
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    d = docx.Document()
    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    gold = RGBColor(0xB8, 0x86, 0x2B)
    ink = RGBColor(0x1E, 0x2A, 0x36)

    header = d.sections[0].header.paragraphs[0]
    header.text = f"{company['short_name']}  ·  {label}  ·  {doc.get('quote_number', '')}"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(0x8A, 0x7E, 0x6B)

    title = d.add_heading(doc["title"], level=0)
    for r in title.runs:
        r.font.color.rgb = ink
    sub = d.add_paragraph(doc.get("subtitle", ""))
    for r in sub.runs:
        r.font.color.rgb = RGBColor(0x8A, 0x7E, 0x6B)

    def shade(cell, hex_fill: str) -> None:
        tc = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tc.append(shd)

    for sec in _visible(doc):
        h = d.add_heading(sec["title"], level=1)
        for r in h.runs:
            r.font.color.rgb = gold
        for b in sec["blocks"]:
            t = b["type"]
            if t == "lead":
                p = d.add_paragraph()
                run = p.add_run(b["text"])
                run.font.size = Pt(14)
                run.font.color.rgb = ink
            elif t in ("paragraph", "note"):
                p = d.add_paragraph(b["text"])
                if t == "note":
                    for r in p.runs:
                        r.font.size = Pt(8.5)
                        r.italic = True
            elif t == "callout":
                p = d.add_paragraph()
                run = p.add_run(b["text"])
                run.bold = True
                run.font.color.rgb = {"good": RGBColor(0x2F, 0x7D, 0x6D), "bad": RGBColor(0xB4, 0x44, 0x3C)}.get(
                    b.get("tone", ""), RGBColor(0xB7, 0x79, 0x1F))
            elif t == "bullets":
                for item in b.get("items", []):
                    d.add_paragraph(item, style="List Bullet")
            elif t == "pagebreak":
                d.add_page_break()
            elif t == "kpis" and b.get("items"):
                table = d.add_table(rows=2, cols=len(b["items"]))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for i, item in enumerate(b["items"]):
                    top, bottom = table.rows[0].cells[i], table.rows[1].cells[i]
                    top.text, bottom.text = item["value"], item["label"]
                    top.paragraphs[0].runs[0].bold = True
                    top.paragraphs[0].runs[0].font.size = Pt(13)
                    bottom.paragraphs[0].runs[0].font.size = Pt(8.5)
                    shade(top, "FBF7EF")
                    shade(bottom, "FBF7EF")
                d.add_paragraph()
            elif t == "bars":
                table = d.add_table(rows=0, cols=3)
                table.style = "Table Grid"
                for item in b.get("items", []):
                    cells = table.add_row().cells
                    cells[0].text = item["label"]
                    cells[1].text = item.get("note", "")
                    cells[2].text = str(item.get("display", ""))
                d.add_paragraph()
            elif t == "table" and b.get("rows"):
                table = d.add_table(rows=1, cols=len(b["header"]))
                table.style = "Table Grid"
                for cell, text in zip(table.rows[0].cells, b["header"]):
                    cell.text = text
                    cell.paragraphs[0].runs[0].bold = True
                    shade(cell, "F4ECDD")
                for row in b["rows"]:
                    for cell, text in zip(table.add_row().cells, row):
                        cell.text = text
                d.add_paragraph()
    d.save(path)
    return path
