"""Generate the PDF and DOCX sample requests used to demonstrate document upload."""

from pathlib import Path

import docx
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

OUT = Path(__file__).resolve().parents[1] / "samples"


def make_pdf() -> None:
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(OUT / "06_brightwater_medical_london.pdf"), pagesize=A4,
                            leftMargin=20 * mm, rightMargin=20 * mm, topMargin=20 * mm, bottomMargin=20 * mm)
    body = styles["BodyText"]
    story = [
        Paragraph("Brightwater Medical Group", styles["Title"]),
        Paragraph("14 Wimpole Street, London W1G 9SX, United Kingdom", body),
        Paragraph("VAT Reg No: GB284716391", body),
        Spacer(1, 8),
        Paragraph("Request for Quotation: Clinic workstation and meeting room refresh", styles["Heading2"]),
        Paragraph("RFQ No: BMG-IT-2026-044", body),
        Paragraph("Date of issue: 21 September 2026", body),
        Spacer(1, 6),
        Paragraph("Brightwater Medical Group operates six outpatient clinics across London. We invite quotations for the "
                  "equipment below, delivered duty paid (DDP) to our London head office.", body),
        Spacer(1, 6),
    ]
    rows = [["Item", "Description", "Qty"],
            ["1", "Business laptops, Core i7, 32GB RAM, 1TB SSD", "30"],
            ["2", "27 inch 4K USB-C monitors", "30"],
            ["3", "USB-C docking stations", "30"],
            ["4", "Wi-Fi 6 access points for clinic floors", "18"],
            ["5", "Video conferencing bar for meeting rooms", "3"],
            ["6", "Endpoint protection licences", "120"]]
    t = Table(rows, colWidths=[15 * mm, 120 * mm, 20 * mm])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, "#999999"), ("BACKGROUND", (0, 0), (-1, 0), "#eeeeee")]))
    story += [t, Spacer(1, 8)]
    for line in [
        "Delivery must be completed within 30 days of the purchase order.",
        "All hardware must carry a minimum 36 months warranty with next business day replacement.",
        "Suppliers must comply with UK GDPR and hold ISO 27001 certification.",
        "Prices must be quoted in GBP. Payment will be made within 30 days of a valid invoice.",
        "Quotations will be evaluated on price (40%), service (40%) and delivery (20%).",
        "Quotations must be submitted by 19 October 2026 to procurement@brightwatermedical.co.uk.",
        "Contact: Helen Carter, Head of Procurement",
    ]:
        story.append(Paragraph(line, body))
    doc.build(story)


def make_docx() -> None:
    d = docx.Document()
    d.add_heading("Lionsgate Analytics Pte. Ltd.", level=1)
    d.add_paragraph("80 Robinson Road, #12-01, Singapore 068898")
    d.add_paragraph("UEN: 201912345K")
    d.add_heading("Request for Proposal – Office expansion (Level 12)", level=2)
    d.add_paragraph("RFP Reference: LGA-OPS-2611")
    d.add_paragraph("Issued on: 23 September 2026")
    d.add_paragraph("Lionsgate Analytics is expanding its Singapore office and requires the following IT equipment "
                    "and services, delivered to our Robinson Road office on a DDP basis.")
    table = d.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Description", "Specification", "Qty")):
        cell.text = text
    for desc, spec, qty in [
        ("Business laptops", "Core i5, 16GB RAM, 512GB SSD", "25"),
        ("24 inch monitors", "Full HD, height adjustable", "25"),
        ("Wireless keyboard and mouse sets", "", "25"),
        ("Microsoft 365 Business Standard", "annual subscription", "25"),
        ("Network switch", "24-port PoE, smart managed", "2"),
        ("On-site installation and configuration", "per device", "25"),
    ]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = desc, spec, qty
    for line in [
        "Delivery within 21 days of purchase order.",
        "Minimum 3 years warranty on all hardware.",
        "Please quote in Singapore dollars; payment terms net 30.",
        "Proposals are due by 14 October 2026. Contact: Wei Ling Tan, Operations Manager, ops@lionsgate-analytics.sg",
    ]:
        d.add_paragraph(line)
    d.save(OUT / "07_lionsgate_singapore.docx")


if __name__ == "__main__":
    make_pdf()
    make_docx()
    print("Samples written to", OUT)
