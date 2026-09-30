"""Phase 9 award-rule strategy, the editable report, its editing assistant and exports."""

import io
from pathlib import Path

import docx
import pytest
from pypdf import PdfReader

from app.agents.base import PipelineContext, StageLog
from app.agents.orchestrator import build_agents
from app.db.seed import load_json
from app.rag.stores import knowledge_store
from app.report.assistant import ReportAssistant, intent_model, resolve_section
from app.report.builder import build_document, build_sections, figures
from app.report.export import render_docx, render_pdf

SAMPLES = Path(__file__).resolve().parents[2] / "samples"


def run(name: str) -> PipelineContext:
    path = SAMPLES / name
    from app.services.documents import extract_text

    ctx = PipelineContext(77, f"A-{name[:2]}", extract_text(path.name, path.read_bytes()), load_json("company.json"),
                          source_filename=path.name, source_path=path)
    for agent in build_agents():
        ctx.messages[agent.produces] = agent.run(ctx, StageLog())
    return ctx


@pytest.fixture(scope="module")
def l1():
    return run("01_sahyadri_campus_refresh.txt")


@pytest.fixture(scope="module")
def qcbs():
    return run("09_konkan_university_rfp.docx")


# --------------------------------------------------------------------------- award analysis


def test_l1_award_analysis_and_no_free_bundles(l1):
    strat = l1.messages["strategy"]
    award = strat.award
    assert award.rule == "L1" and award.lowest_total and award.rank >= 1
    assert award.floor_total <= award.our_total + 1
    # Under a lowest-price award, free services earn no marks, so none are bundled.
    assert all(line.bundle is None for line in strat.lines)
    labels = [o.label for o in award.options]
    assert "Price to become L1" in labels and any(label.startswith("MSE match") for label in labels)
    assert any("no evaluation credit" in r for r in award.reasons)
    if not award.target_feasible:
        assert award.recommendation in ("Stay above floor and use the MSE purchase preference", "Unlikely to win on price")


def test_qcbs_award_uses_compliance_based_technical_score(qcbs):
    award = qcbs.messages["strategy"].award
    assert award.rule == "QCBS"
    assert award.technical_score is not None and 40 <= award.technical_score <= 95
    assert award.combined_score is not None and award.best_competitor_combined is not None
    assert any("70:30" in r for r in award.reasons)


# --------------------------------------------------------------------------- report document and assistant


@pytest.fixture()
def doc_and_helper(l1):
    m = l1.messages
    company = l1.company
    doc = build_document(company, m["parsed"], m["strategy"], m["localisation"], m["proposal"], m.get("compliance"))
    kb = knowledge_store()
    helper = ReportAssistant(
        figures=figures(m["parsed"], m["strategy"], m["localisation"], m.get("compliance")),
        regenerate=lambda key: build_sections(company, m["parsed"], m["strategy"], m["localisation"], m["proposal"],
                                              m.get("compliance"), only=key),
        retrieve=lambda q: kb.evidence(q, k=1),
    )
    return doc, helper


def titles(doc):
    return [s["title"] for s in doc["sections"]]


def test_report_document_has_generated_sections(doc_and_helper):
    doc, _ = doc_and_helper
    keys = [s["key"] for s in doc["sections"]]
    assert keys[:2] == ["summary", "award"] and "risks" in keys and "next" in keys
    assert resolve_section(doc, "timeline")["key"] == "delivery"
    assert resolve_section(doc, "the risks section")["key"] == "risks"


def test_assistant_structural_edits_and_undo(doc_and_helper):
    doc, a = doc_and_helper
    r = a.handle(doc, "Rename risks to Key risks")
    assert r.changed and "Key risks" in titles(doc)
    a.handle(doc, "move delivery above key risks")
    t = titles(doc)
    assert t.index("Delivery roadmap") == t.index("Key risks") - 1
    a.handle(doc, "move next steps to the top")
    assert titles(doc)[0] == "Next steps"
    a.handle(doc, "hide the method section")
    assert next(s for s in doc["sections"] if s["key"] == "method")["hidden"]
    a.handle(doc, "delete the award section")
    assert not any(s["key"] == "award" for s in doc["sections"])
    assert a.handle(doc, "undo").changed and any(s["key"] == "award" for s in doc["sections"])
    assert a.handle(doc, "redo").changed and not any(s["key"] == "award" for s in doc["sections"])
    assert a.handle(doc, "change the title to Bid decision").changed and doc["title"] == "Bid decision"
    assert doc["edited"]


def test_assistant_content_edits(doc_and_helper):
    doc, a = doc_and_helper
    r = a.handle(doc, "add a next step: confirm OEM authorisation letters by Friday")
    assert r.changed and any("Confirm OEM authorisation letters by Friday." in b.get("items", [])
                             for s in doc["sections"] if s["key"] == "next" for b in s["blocks"])
    r = a.handle(doc, "add the total price to the summary")
    assert r.changed and "Total offer value" in r.reply
    r = a.handle(doc, "mention our ISO 27001 certification in risks")
    assert r.changed and "27001" in r.reply
    r = a.handle(doc, "replace 'Next steps' with 'Actions'")
    assert r.changed and "Actions" in titles(doc)
    risks = next(s for s in doc["sections"] if s["key"] == "risks")
    words = sum(len(" ".join(b.get("items", []) + [b.get("text", "")]).split()) for b in risks["blocks"])
    r = a.handle(doc, "shorten the risks section")
    after = sum(len(" ".join(b.get("items", []) + [b.get("text", "")]).split()) for b in risks["blocks"])
    assert not r.changed or after < words
    r = a.handle(doc, "remove the bullet about OEM authorisation")
    assert r.changed
    r = a.handle(doc, "add a section called Client relationship with We have supplied this institute before.")
    assert r.changed and "Client relationship" in titles(doc)
    r = a.handle(doc, "refresh the summary")
    assert r.changed and r.intent == "refresh"


def test_assistant_explains_what_it_cannot_do(doc_and_helper):
    doc, a = doc_and_helper
    r = a.handle(doc, "rename the flux capacitor to Delorean")
    assert not r.changed and "couldn't tell which section" in r.reply
    r = a.handle(doc, "could you perhaps reorder the delivery bit so it comes earlier")
    assert not r.changed and r.intent in ("move_section", "unknown")
    assert a.handle(doc, "replace 'zzz-not-there' with 'x'").changed is False
    assert "rename" in a.handle(doc, "help").reply.lower()


def test_intent_model_generalises():
    model = intent_model()
    preds = model.predict(["please get rid of the risks part", "put delivery before risks", "undo that please",
                           "show me what you can do"])
    assert list(preds) == ["delete_section", "move_section", "undo", "help"]


def test_exports_render(doc_and_helper, tmp_path):
    doc, a = doc_and_helper
    a.handle(doc, "rename risks to Key risks")
    a.handle(doc, "hide the method section")
    company = load_json("company.json")
    pdf = render_pdf(doc, tmp_path / "r.pdf", company)
    text = "\n".join(p.extract_text() for p in PdfReader(pdf).pages)
    assert "K E Y" in text.upper() or "KEY RISKS" in text.upper().replace(" ", " ")
    assert "How these figures" not in text  # hidden sections are left out
    word = docx.Document(render_docx(doc, tmp_path / "r.docx", company))
    headings = [p.text for p in word.paragraphs if p.style.name.startswith("Heading")]
    assert "Key risks" in headings and "How these figures were produced" not in headings


# --------------------------------------------------------------------------- API


def test_report_api_round_trip():
    from fastapi.testclient import TestClient

    from app.main import app
    from tests.test_api import wait_for

    with TestClient(app) as client:
        text = (SAMPLES / "02_kaveri_network_upgrade.txt").read_text()
        rid = client.post("/api/rfps", json={"text": text}).json()["id"]
        wait_for(client, rid)
        doc = client.get(f"/api/rfps/{rid}/report").json()
        assert doc["sections"] and not doc["edited"] and not doc["can_undo"]
        out = client.post(f"/api/rfps/{rid}/report/assistant", json={"message": "rename risks to Key risks"}).json()
        assert out["changed"] and out["document"]["can_undo"] and out["document"]["chat"][-1]["role"] == "assistant"
        sections = out["document"]["sections"]
        sections[0]["blocks"].append({"type": "paragraph", "text": "Typed by the reviewer."})
        saved = client.put(f"/api/rfps/{rid}/report", json={"title": "Kaveri — bid decision", "sections": sections}).json()
        assert saved["title"] == "Kaveri — bid decision" and saved["edited"]
        assert all(b.get("id") for s in saved["sections"] for b in s["blocks"])
        undone = client.post(f"/api/rfps/{rid}/report/undo").json()
        assert undone["title"] != "Kaveri — bid decision" and undone["can_redo"]
        pdf = client.get(f"/api/rfps/{rid}/report/export", params={"format": "pdf"})
        assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf"
        word = client.get(f"/api/rfps/{rid}/report/export", params={"format": "docx"})
        assert word.status_code == 200 and docx.Document(io.BytesIO(word.content)).paragraphs
        edited = client.get(f"/api/rfps/{rid}/documents/report")
        assert "Key risks".upper().replace(" ", "") in "".join(p.extract_text() for p in PdfReader(io.BytesIO(edited.content)).pages).upper().replace(" ", "").replace(" ", "")
