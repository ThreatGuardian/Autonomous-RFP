"""Build the long sample tenders: a 14-page municipal PDF and a university RFP in DOCX.

Run from the repository root: ``python scripts/make_tenders.py``.
"""

from __future__ import annotations

import sys
from pathlib import Path

import docx
from docx.enum.text import WD_BREAK
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

sys.path.insert(0, str(Path(__file__).resolve().parent))
from tender_content import GODAVARI, GV_ORG, GV_REF, KONKAN  # noqa: E402

OUT = Path(__file__).resolve().parents[1] / "samples"

BODY = ParagraphStyle("body", fontName="Times-Roman", fontSize=11.5, leading=16, alignment=TA_JUSTIFY, spaceAfter=5)
CELL = ParagraphStyle("cell", fontName="Times-Roman", fontSize=10.5, leading=13.5)
CELL_B = ParagraphStyle("cellb", parent=CELL, fontName="Times-Bold")
TITLE = ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=16, leading=21, alignment=TA_CENTER, spaceAfter=8, spaceBefore=6)
H1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=12.5, leading=16, spaceBefore=12, spaceAfter=7)
H2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=11, leading=14, spaceBefore=8, spaceAfter=4)


def _decorate(canvas, doc) -> None:
    canvas.saveState()
    w, h = A4
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#333333"))
    canvas.drawString(20 * mm, h - 12 * mm, GV_ORG)
    canvas.drawRightString(w - 20 * mm, h - 12 * mm, f"Tender No. {GV_REF}")
    canvas.setStrokeColor(colors.HexColor("#999999"))
    canvas.line(20 * mm, h - 14 * mm, w - 20 * mm, h - 14 * mm)
    canvas.drawString(20 * mm, 11 * mm, "Signature of bidder with seal")
    canvas.drawRightString(w - 20 * mm, 11 * mm, f"Page {doc.page} of {doc._total}")
    canvas.restoreState()


class _Doc(SimpleDocTemplate):
    _total = 0


def _pdf_story(elements: list[tuple]) -> list:
    story: list = []
    for el in elements:
        kind = el[0]
        if kind == "title":
            story.append(Paragraph(el[1], TITLE))
        elif kind == "h1":
            story.append(Paragraph(el[1], H1))
        elif kind == "h2":
            story.append(Paragraph(el[1], H2))
        elif kind == "p":
            story.append(Paragraph(el[1], BODY))
        elif kind == "clause":
            story.append(Paragraph(f"{el[1]}&nbsp;&nbsp;{el[2]}", ParagraphStyle("c", parent=BODY, leftIndent=11 * mm,
                                                                                    firstLineIndent=-11 * mm)))
        elif kind == "table":
            header, rows, widths = el[1], el[2], el[3]
            data = [[Paragraph(h, CELL_B) for h in header]] + [[Paragraph(c, CELL) for c in r] for r in rows]
            t = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#666666")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6e6e6")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story += [t, Spacer(1, 7)]
        elif kind == "break":
            story.append(PageBreak())
    # Keep each specification heading with its table.
    out: list = []
    for item in story:
        if out and isinstance(item, Table) and isinstance(out[-1], Paragraph) and out[-1].style.name == "h2":
            out[-1] = KeepTogether([out[-1], item])
        else:
            out.append(item)
    return out


def make_pdf(path: Path) -> int:
    def build(total: int) -> int:
        doc = _Doc(str(path), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=22 * mm, bottomMargin=20 * mm,
                   title=f"Tender {GV_REF}", author=GV_ORG)
        doc._total = total
        doc.build(_pdf_story(GODAVARI), onFirstPage=_decorate, onLaterPages=_decorate)
        return doc.page

    return build(build(0))  # second pass prints the correct "of N"


def make_docx(path: Path) -> None:
    d = docx.Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    for el in KONKAN:
        kind = el[0]
        if kind == "title":
            d.add_heading(el[1], level=0)
        elif kind == "h1":
            d.add_heading(el[1], level=1)
        elif kind == "h2":
            d.add_heading(el[1], level=2)
        elif kind == "p":
            d.add_paragraph(el[1])
        elif kind == "clause":
            d.add_paragraph(f"{el[1]}\t{el[2]}")
        elif kind == "table":
            header, rows = el[1], el[2]
            t = d.add_table(rows=1, cols=len(header))
            t.style = "Table Grid"
            for cell, text in zip(t.rows[0].cells, header):
                cell.text = text
                for r in cell.paragraphs[0].runs:
                    r.bold = True
            for row in rows:
                cells = t.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = text
            d.add_paragraph()
        elif kind == "break":
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.save(path)


if __name__ == "__main__":
    pages = make_pdf(OUT / "08_godavari_smart_city_tender.pdf")
    make_docx(OUT / "09_konkan_university_rfp.docx")
    print(f"Tender PDF: {pages} pages; university RFP written to {OUT}")
