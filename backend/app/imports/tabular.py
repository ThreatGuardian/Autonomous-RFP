"""Read CSV and Excel sheets exported by accounting and ERP tools, and map their headers.

SME data rarely arrives with our column names: a Tally or Busy export says
"Item Name", "Stock Group", "Closing Qty" and "Purchase Rate"; a distributor price
list says "Part No.", "Make" and "Dealer Price". Headers are matched to fields
through synonym lists after normalisation, and values are parsed with Indian
number conventions (lakh grouping, "Rs."/"₹", "18%", "3 years").
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from typing import Any


def norm(header: str) -> str:
    h = header.lower().replace("&", " and ")
    h = re.sub(r"[^a-z0-9%]+", " ", h)
    return re.sub(r"\s+", " ", h).strip()


@dataclass
class Table:
    headers: list[str]
    rows: list[dict[str, str]]
    sheet: str | None = None
    header_row: int = 1  # 1-based row number of the header in the source


def _rows_to_table(raw: list[list[Any]], sheet: str | None = None) -> Table:
    cells = [["" if c is None else str(c).strip() for c in row] for row in raw]
    # Header = the first of the top ten rows with the most non-numeric text cells
    # (exports often start with a company name and a period line).
    best, best_score = 0, -1
    for i, row in enumerate(cells[:10]):
        score = sum(1 for c in row if c and not re.fullmatch(r"[-\d.,₹% ]+", c))
        if score > best_score:
            best, best_score = i, score
    headers = [h or f"Column {j + 1}" for j, h in enumerate(cells[best] if cells else [])]
    rows = []
    for row in cells[best + 1:]:
        if not any(row):
            continue
        rows.append({headers[j]: (row[j] if j < len(row) else "") for j in range(len(headers))})
    return Table(headers, rows, sheet, best + 1)


def read_table(filename: str, data: bytes) -> Table:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return _rows_to_table([list(r) for r in ws.iter_rows(values_only=True)], ws.title)
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return _rows_to_table(list(csv.reader(io.StringIO(text), dialect)))


@dataclass
class Mapping:
    fields: dict[str, str] = field(default_factory=dict)  # field -> source header
    unmapped: list[str] = field(default_factory=list)

    def get(self, row: dict[str, str], name: str) -> str:
        header = self.fields.get(name)
        return (row.get(header) or "").strip() if header else ""


def map_headers(headers: list[str], synonyms: dict[str, list[str]]) -> Mapping:
    """Assign each field the best matching header: exact synonym, then all-words match."""
    mapping = Mapping()
    taken: set[str] = set()
    normed = {h: norm(h) for h in headers}
    for exact in (True, False):
        for fld, words in synonyms.items():
            if fld in mapping.fields:
                continue
            for syn in words:
                for h, n in normed.items():
                    if h in taken:
                        continue
                    if (n == syn) if exact else (set(syn.split()) <= set(n.split())):
                        mapping.fields[fld] = h
                        taken.add(h)
                        break
                if fld in mapping.fields:
                    break
    mapping.unmapped = [h for h in headers if h not in taken]
    return mapping


# --------------------------------------------------------------------------- values

_NUM = re.compile(r"-?\d[\d,]*(?:\.\d+)?")


def parse_number(value: str) -> float | None:
    """'₹1,23,456.50' -> 123456.5, '45,000/Nos' -> 45000, '2.5 lakh' -> 250000."""
    if value is None:
        return None
    v = str(value).strip().lower()
    if not v or v in ("-", "na", "n/a", "nil"):
        return None
    m = _NUM.search(v.replace(" ", "") if re.fullmatch(r"[\d ,.]+", v) else v)
    if not m:
        return None
    num = float(m.group(0).replace(",", ""))
    if re.search(r"\blakhs?\b|\blacs?\b", v):
        num *= 1e5
    elif re.search(r"\bcrores?\b|\bcr\b", v):
        num *= 1e7
    return num


def parse_months(value: str) -> int | None:
    v = str(value or "").lower()
    n = parse_number(v)
    if n is None:
        return 0 if re.search(r"\bno warranty\b|\bnil\b", v) else None
    if re.search(r"\by(?:ea)?rs?\b|\byr\b", v):
        return int(round(n * 12))
    if re.search(r"\blifetime\b", v):
        return 120
    return int(round(n))


def parse_int(value: str) -> int | None:
    n = parse_number(value)
    return None if n is None else int(round(n))
