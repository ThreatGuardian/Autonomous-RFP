from pathlib import Path

import pytest

from app.agents.base import PipelineContext, StageLog
from app.agents.parser_agent import RfpParserAgent
from app.db.seed import load_json
from app.nlp.extractors import extract_currency, extract_dates, extract_location, extract_terms
from app.nlp.line_items import build_item
from app.services.documents import extract_text

SAMPLES = Path(__file__).resolve().parents[2] / "samples"


def parse(name: str):
    text = (SAMPLES / name).read_text()
    return RfpParserAgent().run(PipelineContext(1, "T", text, load_json("company.json")), StageLog())


@pytest.mark.parametrize(
    "text,qty,desc",
    [
        ("25 x Dell Latitude laptops, 16GB RAM", 25, "Dell Latitude laptops, 16GB RAM"),
        ("- 48-port PoE+ switches: Qty 4", 4, "48-port PoE+ switches"),
        ("Supply of forty 24-inch monitors", 40, "24-inch monitors"),
        ("Wi-Fi 6 access points (30 units)", 30, "Wi-Fi 6 access points"),
        ("3. Online UPS 3kVA - 6", 6, "Online UPS 3kVA"),
        ("- 40 x 27 inch 4K monitors", 40, "27 inch 4K monitors"),
    ],
)
def test_quantity_extraction(text, qty, desc):
    item = build_item(text)
    assert item is not None and item.quantity == qty and item.description == desc


@pytest.mark.parametrize("text", ["Laptop with 3 years warranty", "RFQ No: HF-PROC-2611", "16GB RAM minimum"])
def test_spec_numbers_are_not_quantities(text):
    item = build_item(text)
    assert item is None or item.quantity not in (3, 2611, 16)


def test_specs_extracted():
    item = build_item("10 x Rack server, 2x Xeon Silver, 64GB RAM, 1.92TB SSD, 27 inch")
    assert item.specs["ram_gb"] == 64 and item.specs["storage_gb"] == 1920 and item.specs["cpu_family"] == "xeon"


def test_location_currency_dates():
    text = "Deliver to: Fresno, California\nPrices should be quoted in US dollars.\nProposals are due by 10/20/2026."
    loc = extract_location(text)
    assert (loc.country, loc.region) == ("US", "California")
    assert extract_currency(text, "US")[0] == "USD"
    assert extract_dates(text, month_first=True)["due"] == "2026-10-20"


def test_terms():
    t = extract_terms("Delivery within six weeks on a DDP basis. Payment: net 45. Minimum 3 years warranty. 30% advance.")
    assert t["incoterm"] == "DDP" and t["delivery_days"] == 42 and t["payment_days"] == 45
    assert t["warranty_months_required"] == 36 and t["advance_pct"] == 30


def test_domestic_rfp_end_to_end():
    parsed = parse("01_sahyadri_campus_refresh.txt")
    assert parsed.client.name == "Sahyadri Institute of Technology"
    assert (parsed.client.country, parsed.client.region, parsed.client.segment) == ("IN", "Maharashtra", "education")
    assert parsed.currency.code == "INR" and parsed.due_date == "2026-10-10"
    assert [i.quantity for i in parsed.line_items] == [60, 120, 120, 4, 24, 4, 30, 120]
    assert all(i.status == "matched" for i in parsed.line_items)
    assert parsed.terms.lowest_price_award


def test_table_rfp():
    parsed = parse("02_kaveri_network_upgrade.txt")
    skus = [i.selected_sku for i in parsed.line_items]
    assert skus == ["MSS-NW-501", "MSS-NW-506", "MSS-SV-601", "MSS-ST-701", "MSS-RK-901", "MSS-SW-123", "MSS-SV-132"]
    assert parsed.terms.price_weight_pct == 60


def test_international_rfp_and_unmatched_item():
    parsed = parse("04_cedar_valley_school.txt")
    assert (parsed.client.country, parsed.currency.code, parsed.terms.incoterm) == ("US", "USD", "DDP")
    whiteboard = parsed.line_items[-1]
    assert whiteboard.status == "unmatched" and whiteboard.selected_sku is None
    assert any("no suitable catalogue match" in w for w in parsed.warnings)


def test_eu_rfp():
    parsed = parse("05_rheinland_engineering.txt")
    assert parsed.client.is_eu and parsed.currency.code == "EUR" and parsed.terms.incoterm == "DAP"
    assert parsed.client_reference == "RPT-EK-2026-118" and parsed.due_date == "2026-10-16"


def test_docx_ingestion():
    import io

    import docx

    doc = docx.Document()
    doc.add_paragraph("Request for Quotation")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Description", "Qty"
    table.cell(1, 0).text, table.cell(1, 1).text = "27 inch monitors", "12"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text("rfq.docx", buf.getvalue())
    assert "| 27 inch monitors | 12 |" in text
