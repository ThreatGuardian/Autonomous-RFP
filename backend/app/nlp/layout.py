"""Layout analysis for long tender documents.

A 10–15 page tender is not a bag of sentences: it has numbered sections,
running headers and footers, key-data tables and annexure forms. This module
turns PDF, DOCX and plain text into an ordered list of :class:`Block` objects
that keep the signals later stages need:

* the page each block sits on (real pages for PDF, rendered page breaks for
  DOCX, form-feeds or an estimate for text);
* typographic emphasis (larger or bold type, Word heading styles) so section
  headings can be told apart from numbered clauses;
* table rows as cells, so specification and schedule tables stay tabular;
* repeated running headers and footers ("Tender No. … Page 3 of 14"),
  removed so they are not read as requirements.

PDFs are read with PyMuPDF when it is installed (fonts, positions and ruled
tables) and with pypdf otherwise. Scanned pages without a text layer are
reported; they are OCR'd only when Tesseract is available on the host.
"""

from __future__ import annotations

import io
import re
from collections import Counter
from dataclasses import asdict, dataclass, field
from typing import Any

CHARS_PER_PAGE = 3200  # used only to estimate pages for text without page breaks


@dataclass
class Block:
    kind: str  # "text" | "row"
    text: str
    page: int
    style: str = "body"  # "body" | "heading" | "title"
    level: int | None = None  # heading level when the source declares one (Word styles)
    cells: list[str] | None = None
    table: int | None = None  # table id shared by the rows of one table

    def as_dict(self) -> dict[str, Any]:
        return {k: v for k, v in asdict(self).items() if v is not None}


@dataclass
class Layout:
    format: str
    blocks: list[Block]
    pages: int
    pages_estimated: bool = False
    scanned_pages: list[int] = field(default_factory=list)
    removed_lines: int = 0
    tables: int = 0
    notes: list[str] = field(default_factory=list)

    def to_text(self) -> str:
        """Plain text with tables as ``| a | b |`` rows and form-feeds between pages."""
        out: list[str] = []
        page = 1
        for b in self.blocks:
            while b.page > page:
                out.append("\f")
                page += 1
            out.append("| " + " | ".join(b.cells or []) + " |" if b.kind == "row" else b.text)
        return "\n".join(out).replace("\n\f\n", "\n\f")

    def stats(self) -> dict[str, Any]:
        return {"format": self.format, "pages": self.pages, "pages_estimated": self.pages_estimated,
                "blocks": len(self.blocks), "tables": self.tables, "scanned_pages": self.scanned_pages,
                "removed_header_footer_lines": self.removed_lines}


# --------------------------------------------------------------------------- shared helpers

_PAGE_NO = re.compile(r"^\s*(?:page\s*)?\d{1,3}\s*(?:(?:of|/)\s*\d{1,3})?\s*$|\bpage\s+\d{1,3}\s*(?:of|/)\s*\d{1,3}\b", re.I)


def _shape(text: str) -> str:
    """Page-independent form of a line: digits collapsed so 'Page 3 of 14' == 'Page 4 of 14'."""
    return re.sub(r"\d+", "#", re.sub(r"\s+", " ", text.strip().lower()))


def _dehyphenate(lines: list[str]) -> str:
    out = ""
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if out.endswith("-") and line[:1].islower():
            out = out[:-1] + line
        else:
            out = f"{out} {line}" if out else line
    return out


# --------------------------------------------------------------------------- PDF


def analyse_pdf(data: bytes) -> Layout:
    try:
        import pymupdf  # noqa: F401
    except ImportError:  # pragma: no cover - exercised only without PyMuPDF
        return _pdf_fallback(data)
    return _pdf_pymupdf(data)


def _pdf_pymupdf(data: bytes) -> Layout:
    import pymupdf

    if hasattr(pymupdf, "no_recommend_layout"):
        pymupdf.no_recommend_layout()
    doc = pymupdf.open(stream=data, filetype="pdf")
    raw: list[dict[str, Any]] = []  # candidate lines with geometry, before header/footer removal
    size_chars: Counter[float] = Counter()
    scanned: list[int] = []
    table_count = 0
    notes: list[str] = []
    for pno, page in enumerate(doc, start=1):
        height = page.rect.height or 842
        tables = []
        try:
            for t in page.find_tables().tables:
                rows = [[re.sub(r"\s+", " ", c or "").strip() for c in r] for r in t.extract()]
                rows = [r for r in rows if any(r)]
                if len(rows) >= 2 and max(len(r) for r in rows) >= 2:
                    tables.append((pymupdf.Rect(t.bbox), rows))
        except Exception:  # table finder is best-effort
            tables = []
        page_dict = page.get_text("dict")
        text_blocks = [b for b in page_dict["blocks"] if b.get("type") == 0]
        has_text = any(s["text"].strip() for b in text_blocks for l in b["lines"] for s in l["spans"])
        if not has_text:
            if any(b.get("type") == 1 for b in page_dict["blocks"]) or page.get_images():
                ocr = _ocr_page(page)
                if ocr:
                    notes.append(f"Page {pno} was scanned and has been read by OCR.")
                    for para in re.split(r"\n\s*\n", ocr):
                        if para.strip():
                            raw.append({"page": pno, "y": 0, "text": _dehyphenate(para.splitlines()), "size": 0, "bold": False,
                                        "margin": False, "kind": "text"})
                else:
                    scanned.append(pno)
            continue
        for rect, rows in tables:
            table_count += 1
            for r in rows:
                raw.append({"page": pno, "y": rect.y0, "kind": "row", "cells": r, "table": table_count, "margin": False})
        for b in text_blocks:
            group: list[dict[str, Any]] = []
            for line in b["lines"]:
                spans = [s for s in line["spans"] if s["text"].strip()]
                if not spans:
                    continue
                bbox = pymupdf.Rect(line["bbox"])
                if any(rect.intersects(bbox) and rect.contains(bbox.tl + (1, 1)) for rect, _ in tables):
                    continue
                # Wide gaps between spans are column gaps of rule-less tables: keep them as tabs.
                text = spans[0]["text"]
                for prev, cur in zip(spans, spans[1:]):
                    gap = cur["bbox"][0] - prev["bbox"][2]
                    text += ("\t" if gap > 2.2 * max(cur["size"], 6) else ("" if gap < 0.8 else " ")) + cur["text"]
                size = round(max(s["size"] for s in spans), 1)
                chars = sum(len(s["text"].strip()) for s in spans)
                bold = all((s["flags"] & 16) or "bold" in s["font"].lower() or "semibold" in s["font"].lower() for s in spans)
                size_chars[size] += chars
                margin = bbox.y1 < height * 0.075 or bbox.y0 > height * 0.925
                group.append({"page": pno, "y": bbox.y0, "text": text.strip(), "size": size, "bold": bold, "margin": margin,
                              "kind": "text"})
            # Split the PyMuPDF block where the typography changes (heading directly above its body).
            current: list[dict[str, Any]] = []
            for ln in group:
                if current and (ln["size"] != current[-1]["size"] or ln["bold"] != current[-1]["bold"] or ln["margin"]
                                or current[-1]["margin"] or "\t" in ln["text"] or "\t" in current[-1]["text"]):
                    raw.append(_merge(current))
                    current = []
                current.append(ln)
            if current:
                raw.append(_merge(current))
    pages = doc.page_count
    raw.sort(key=lambda r: (r["page"], r["y"]))

    # Running headers and footers: margin lines whose shape repeats on many pages.
    shapes = Counter(_shape(r["text"]) for r in raw if r["kind"] == "text" and r["margin"])
    threshold = max(2, int(pages * 0.4))
    repeated = {s for s, n in shapes.items() if n >= threshold}
    body = size_chars.most_common(1)[0][0] if size_chars else 10.0
    blocks: list[Block] = []
    removed = 0
    headers: dict[int, list[str]] = {}
    alias: dict[int, int] = {}
    for r in raw:
        if r["kind"] == "row":
            tid = r["table"]
            if tid not in headers:
                # A table continued on the next page repeats its header row: stitch it to the previous table.
                prev = next((b for b in reversed(blocks) if b.kind == "text" or b.kind == "row"), None)
                if prev is not None and prev.kind == "row" and headers.get(prev.table) == r["cells"]:
                    alias[tid] = prev.table
                    headers[tid] = r["cells"]
                    continue
                headers[tid] = r["cells"]
            r = {**r, "table": alias.get(tid, tid)}
            blocks.append(Block("row", " | ".join(c for c in r["cells"] if c), r["page"], cells=r["cells"], table=r["table"]))
            continue
        if r["margin"] and (_shape(r["text"]) in repeated or _PAGE_NO.search(r["text"])):
            removed += 1
            continue
        style = "body"
        words = len(r["text"].split())
        if r["size"] >= body * 1.45 and words <= 16:
            style = "title"
        elif (r["size"] >= body * 1.12 or r["bold"]) and words <= 18 and not r["text"].rstrip().endswith((".", ";", ",")):
            style = "heading"
        if "\t" in r["text"]:
            cells = [c.strip() for c in r["text"].split("\t")]
            blocks.append(Block("text", r["text"], r["page"], style=style, cells=cells if len(cells) >= 2 else None))
        else:
            blocks.append(Block("text", r["text"], r["page"], style=style))
    if scanned:
        notes.append(f"{len(scanned)} scanned page(s) without a text layer could not be read (pages "
                     f"{', '.join(map(str, scanned[:8]))}); install Tesseract OCR to read them.")
    return Layout("pdf", blocks, pages, scanned_pages=scanned, removed_lines=removed, tables=table_count, notes=notes)


def _merge(lines: list[dict[str, Any]]) -> dict[str, Any]:
    first = dict(lines[0])
    first["text"] = _dehyphenate([l["text"] for l in lines]) if "\t" not in first["text"] else first["text"]
    first["margin"] = all(l["margin"] for l in lines)
    return first


def _ocr_page(page) -> str | None:
    """OCR one page when Tesseract is installed; otherwise ``None``."""
    try:
        import shutil

        if not shutil.which("tesseract"):
            return None
        tp = page.get_textpage_ocr(dpi=300, full=True)
        return page.get_text("text", textpage=tp)
    except Exception:
        return None


def _pdf_fallback(data: bytes) -> Layout:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text(extraction_mode="layout") or "")
        except TypeError:
            pages.append(page.extract_text() or "")
    layout = analyse_text("\f".join(pages))
    layout.format = "pdf"
    layout.pages_estimated = False
    layout.scanned_pages = [i for i, p in enumerate(pages, start=1) if not p.strip()]
    return layout


# --------------------------------------------------------------------------- DOCX

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def analyse_docx(data: bytes) -> Layout:
    import docx
    from docx.table import Table

    document = docx.Document(io.BytesIO(data))
    blocks: list[Block] = []
    page = 1
    explicit_breaks = False
    chars_on_page = 0
    table_id = 0
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            breaks = sum(1 for n in child.iter() if (n.tag == f"{_W}br" and n.get(f"{_W}type") == "page")
                         or n.tag == f"{_W}lastRenderedPageBreak")
            text = "".join((n.text or "") if n.tag == f"{_W}t" else " " for n in child.iter()
                           if n.tag in (f"{_W}t", f"{_W}tab")).strip()
            if breaks and not text:
                explicit_breaks = True
                page += breaks
                chars_on_page = 0
                continue
            if breaks:
                explicit_breaks = True
                page += breaks
                chars_on_page = 0
            if not text:
                continue
            style_name = ""
            p_style = child.find(f"{_W}pPr/{_W}pStyle")
            if p_style is not None:
                sid = p_style.get(f"{_W}val") or ""
                try:
                    style_name = document.styles.get_by_id(sid).name or sid
                except Exception:
                    style_name = sid
            style, level = "body", None
            m = re.match(r"heading\s*(\d)", style_name, re.I)
            if m:
                style, level = "heading", int(m.group(1))
            elif style_name.lower() in {"title", "subtitle"}:
                style = "title"
            else:
                runs = [r for r in child.iter(f"{_W}r") if "".join(t.text or "" for t in r.iter(f"{_W}t")).strip()]
                if runs and all(r.find(f"{_W}rPr/{_W}b") is not None for r in runs) and len(text.split()) <= 18 \
                        and not text.endswith((".", ";", ",")):
                    style = "heading"
            if not explicit_breaks and chars_on_page > CHARS_PER_PAGE:
                page += 1
                chars_on_page = 0
            chars_on_page += len(text)
            blocks.append(Block("text", text, page, style=style, level=level))
        elif tag == "tbl":
            table_id += 1
            table = Table(child, document)
            for row in table.rows:
                cells: list[str] = []
                for c in row.cells:
                    value = re.sub(r"\s+", " ", c.text).strip()
                    if not cells or value != cells[-1] or not value:  # merged cells repeat their text
                        cells.append(value)
                if any(cells):
                    blocks.append(Block("row", " | ".join(c for c in cells if c), page, cells=cells, table=table_id))
                    chars_on_page += sum(len(c) for c in cells)
        elif tag == "sectPr":
            continue
    return Layout("docx", blocks, page, pages_estimated=not explicit_breaks, tables=table_id)


# --------------------------------------------------------------------------- text

_TABLE_ROW = re.compile(r"^\s*\|.*\|\s*$")


def analyse_text(text: str) -> Layout:
    from app.nlp.text import normalize

    raw_pages = text.split("\f")
    explicit = len(raw_pages) > 1
    blocks: list[Block] = []
    table_id = 0
    in_table = False
    page_no = 0
    for raw_page in raw_pages:
        page_no += 1
        chars = 0
        for para in re.split(r"\n", normalize(raw_page)):
            line = para.strip()
            if not line:
                in_table = False
                continue
            if not explicit and chars > CHARS_PER_PAGE:
                page_no += 1
                chars = 0
            chars += len(line)
            if _TABLE_ROW.match(line) or "\t" in line:
                if re.fullmatch(r"[\s|:\-+=]+", line):
                    continue
                cells = [c.strip() for c in (line.strip().strip("|").split("|") if "|" in line else line.split("\t"))]
                if sum(1 for c in cells if c) >= 2:
                    if not in_table:
                        table_id += 1
                        in_table = True
                    blocks.append(Block("row", " | ".join(c for c in cells if c), page_no, cells=cells, table=table_id))
                    continue
            in_table = False
            blocks.append(Block("text", line, page_no))
    # Drop running headers/footers of text extracted from paginated sources.
    removed = 0
    if explicit and page_no >= 3:
        firsts = Counter()
        for p in range(1, page_no + 1):
            on_page = [b for b in blocks if b.page == p and b.kind == "text"]
            for b in on_page[:2] + on_page[-2:]:
                firsts[_shape(b.text)] += 1
        repeated = {s for s, n in firsts.items() if n >= max(2, int(page_no * 0.4))}
        kept = []
        for b in blocks:
            if b.kind == "text" and (_shape(b.text) in repeated or _PAGE_NO.fullmatch(b.text.strip())):
                removed += 1
                continue
            kept.append(b)
        blocks = kept
    return Layout("text", blocks, max(page_no, 1), pages_estimated=not explicit, removed_lines=removed, tables=table_id)


def analyse(filename: str | None, data: bytes | None, text: str) -> Layout:
    """Best available layout: the original file when kept, otherwise the stored text."""
    suffix = (filename or "").rsplit(".", 1)[-1].lower() if filename and "." in filename else ""
    if data is not None:
        try:
            if suffix == "pdf":
                return analyse_pdf(data)
            if suffix == "docx":
                return analyse_docx(data)
        except Exception:  # corrupt original: fall back to the extracted text
            pass
    layout = analyse_text(text)
    if suffix in {"pdf", "docx"}:
        layout.format = suffix
    return layout

