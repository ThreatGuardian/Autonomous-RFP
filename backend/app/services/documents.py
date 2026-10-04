"""Text extraction from uploaded RFP documents (PDF, DOCX, plain text)."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path


class UnsupportedDocument(ValueError):
    pass


SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".text"}
MAX_DOCX_EXPANDED_BYTES = 80 * 1024 * 1024
MAX_DOCX_ENTRIES = 3000


def _check_pdf(data: bytes) -> None:
    if not data.lstrip()[:5].startswith(b"%PDF-"):
        raise UnsupportedDocument("The file is not a valid PDF.")


def _check_docx(data: bytes) -> None:
    """Reject files that are not Word documents, and archives that expand to an unreasonable size."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
    except zipfile.BadZipFile:
        raise UnsupportedDocument("The file is not a valid Word document.") from None
    if len(infos) > MAX_DOCX_ENTRIES or sum(i.file_size for i in infos) > MAX_DOCX_EXPANDED_BYTES:
        raise UnsupportedDocument("The Word document is too large to process.")
    if not any(i.filename == "word/document.xml" for i in infos):
        raise UnsupportedDocument("The file is not a valid Word document.")


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        _check_pdf(data)
        return _pdf(data)
    if suffix == ".docx":
        _check_docx(data)
        return _docx(data)
    if suffix in {".txt", ".md", ".text", ""}:
        for enc in ("utf-8", "utf-16", "cp1252", "latin-1"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
    raise UnsupportedDocument(f"Unsupported file type '{suffix or filename}'. Upload PDF, DOCX or TXT.")


def _pdf(data: bytes) -> str:
    from app.nlp.layout import analyse_pdf

    try:
        layout = analyse_pdf(data)
    except ValueError as exc:
        raise UnsupportedDocument(str(exc)) from None
    except RuntimeError:  # PyMuPDF raises RuntimeError subclasses for damaged or encrypted files
        raise UnsupportedDocument("The PDF could not be read; it may be damaged or password-protected.") from None
    text = layout.to_text().strip()
    if not text.replace("\f", "").strip():
        raise UnsupportedDocument("The PDF contains no extractable text (it may be a scanned image without OCR).")
    return text


def _pdf_plain(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text(extraction_mode="layout") or "")
        except TypeError:  # older pypdf without layout mode
            pages.append(page.extract_text() or "")
    text = "\n\n".join(pages).strip()
    if not text:
        raise UnsupportedDocument("The PDF contains no extractable text (it may be a scanned image).")
    return text


def _docx(data: bytes) -> str:
    from app.nlp.layout import analyse_docx

    return analyse_docx(data).to_text()


def _docx_plain(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts: list[str] = []
    # Preserve body order of paragraphs and tables.
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = "".join(node.text or "" for node in child.iter() if node.tag.endswith("}t"))
            if text.strip():
                parts.append(text)
        elif tag == "tbl":
            table = docx.table.Table(child, document)
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                parts.append("| " + " | ".join(cells) + " |")
            parts.append("")
    return "\n".join(parts)
